# -*- coding: utf-8 -*-
"""把主流程测试文档（Markdown）同时转换为 Word 与 自包含 HTML"""
import re
import os
import html as htmlmod
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SRC = r"D:\教学工作\2026\月工作总结以及工作计划\课时签到管理后台_主流程测试文档.md"
DOCX = r"D:\教学工作\2026\月工作总结以及工作计划\课时签到管理后台_主流程测试文档.docx"
HTML = r"D:\教学工作\2026\月工作总结以及工作计划\课时签到管理后台_主流程测试文档.html"

md = open(SRC, encoding="utf-8").read()
lines = md.split("\n")

# ============================================================ Word
doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.3
for sec in doc.sections:
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)


def sr(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color: run.font.color.rgb = RGBColor(*color)


def rich(p, text, size=10.5, name="宋体", base_bold=False):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            sr(p.add_run(part[2:-2]), name=name, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            sr(p.add_run(part[1:-1]), name="Consolas", size=size - 0.5, bold=base_bold, color=(0xC7, 0x25, 0x4E))
        else:
            sr(p.add_run(part), name=name, size=size, bold=base_bold)


def flush_table(rows):
    if not rows: return
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(cols):
            txt = row[j] if j < len(row) else ""
            cells[j].text = ""
            p = cells[j].paragraphs[0]; p.paragraph_format.space_after = Pt(1)
            rich(p, txt, size=9.5, name="黑体" if i == 0 else "宋体", base_bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


tb = []; incode = False
for raw in lines:
    line = raw.rstrip()
    if line.strip().startswith("```"):
        if incode: incode = False
        else:
            if tb: flush_table(tb); tb = []
            incode = True
        continue
    if incode:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
        sr(p.add_run(line if line else " "), name="Consolas", size=9, color=(0x33, 0x33, 0x33)); continue
    if line.strip().startswith("|"):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c): continue
        tb.append(cells); continue
    elif tb:
        flush_table(tb); tb = []
    s = line.strip()
    if not s or s == "---": continue
    if s.startswith("### "):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
        rich(p, s[4:], size=12, name="黑体", base_bold=True)
    elif s.startswith("## "):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(5)
        rich(p, s[3:], size=14, name="黑体", base_bold=True)
        for r in p.runs: r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    elif s.startswith("# "):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
        rich(p, s[2:], size=18, name="黑体", base_bold=True)
        for r in p.runs: r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    elif re.match(r"^[-*] ", s):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.7); p.paragraph_format.space_after = Pt(2)
        rich(p, "· " + s[2:])
    else:
        p = doc.add_paragraph(); rich(p, s)
if tb: flush_table(tb)
doc.save(DOCX)
print("Word 已生成：" + DOCX)

# ============================================================ HTML
def inline(text):
    t = htmlmod.escape(text)
    t = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", t)
    t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
    return t


out = []
tb = []; incode = False
for raw in lines:
    line = raw.rstrip()
    if line.strip().startswith("```"):
        if incode:
            out.append("</pre>"); incode = False
        else:
            if tb: out.append("<table>" + "".join("<tr>" + "".join("<td>%s</td>" % inline(c) for c in r) + "</tr>" for r in tb) + "</table>"); tb = []
            out.append("<pre>"); incode = True
        continue
    if incode:
        out.append(htmlmod.escape(line)); continue
    if line.strip().startswith("|"):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            if tb:
                head = tb.pop(0)
                out.append("<table><thead><tr>" + "".join("<th>%s</th>" % inline(c) for c in head) + "</tr></thead><tbody>" +
                           "".join("<tr>" + "".join("<td>%s</td>" % inline(c) for c in r) + "</tr>" for r in tb) + "</tbody></table>")
                tb = []
            continue
        tb.append(cells); continue
    elif tb:
        out.append("<table>" + "".join("<tr>" + "".join("<td>%s</td>" % inline(c) for c in r) + "</tr>" for r in tb) + "</table>"); tb = []
    s = line.strip()
    if not s or s == "---":
        continue
    if s.startswith("### "): out.append("<h3>%s</h3>" % inline(s[4:]))
    elif s.startswith("## "): out.append("<h2>%s</h2>" % inline(s[3:]))
    elif s.startswith("# "): out.append("<h1>%s</h1>" % inline(s[2:]))
    elif re.match(r"^[-*] ", s): out.append("<div class='li'>· %s</div>" % inline(s[2:]))
    elif s.startswith(">"): out.append("<blockquote>%s</blockquote>" % inline(s.lstrip("> ").strip()))
    else: out.append("<p>%s</p>" % inline(s))
if tb: out.append("<table>" + "".join("<tr>" + "".join("<td>%s</td>" % inline(c) for c in r) + "</tr>" for r in tb) + "</table>")
if incode: out.append("</pre>")

page = """<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>课时签到管理后台 · 主流程测试文档</title>
<style>
 body{font-family:-apple-system,"Microsoft YaHei",sans-serif;line-height:1.7;color:#24292f;
      max-width:1000px;margin:0 auto;padding:28px 20px 60px;background:#fff}
 h1{font-size:26px;text-align:center;color:#2563eb;margin:10px 0 24px}
 h2{font-size:20px;margin:34px 0 12px;padding-left:10px;border-left:5px solid #2563eb;color:#1f3a5f}
 h3{font-size:16px;margin:22px 0 8px;color:#1f3a5f}
 table{border-collapse:collapse;width:100%;margin:10px 0 16px;font-size:13.5px}
 th,td{border:1px solid #d0d7de;padding:7px 10px;text-align:left;vertical-align:top}
 th{background:#f3f6fb;font-weight:700}
 tr:nth-child(even) td{background:#fafbfc}
 code{background:#f6f8fa;padding:1px 5px;border-radius:4px;color:#c7254e;font-size:13px}
 pre{background:#f6f8fa;border:1px solid #e1e4e8;border-radius:8px;padding:12px 14px;overflow:auto;font-size:13px}
 blockquote{border-left:4px solid #d0d7de;margin:10px 0;padding:4px 12px;color:#57606a}
 .li{margin:3px 0 3px 18px}
 .banner{background:#eff6ff;border:1px solid #bfdbfe;border-radius:10px;padding:12px 16px;margin-bottom:22px;font-size:14px}
</style></head><body>
<div class="banner">课堂课时确认与学时统计 · 教务管理后台（Web 端）主流程测试文档 ·
被测地址 <a href="http://8.138.41.74/admin/">http://8.138.41.74/admin/</a> ·
执行结果 <b>17 / 17 全部通过</b> · 可下载 <a href="test-report.docx">Word 版</a><br>
<span style="color:#b45309">访问说明：本页仅供内部查看，访问需认证（账号与密码请向系统管理员索取），请勿外传。</span></div>
""" + "\n".join(out) + "\n</body></html>"
open(HTML, "w", encoding="utf-8").write(page)
print("HTML 已生成：" + HTML)
print("段落 %d / 表格 %d" % (len(doc.paragraphs), len(doc.tables)))
